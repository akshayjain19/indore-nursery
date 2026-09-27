#!/usr/bin/env python3
"""One-time / refresh migration: legacy catalog.json → Excel source of truth."""

from __future__ import annotations

import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from openpyxl.styles import Font

from lib.config import DEFAULT_ACTIVE_PLANT_SLUGS, EXCEL_PATH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(ROOT, "data", "catalog.json")

PLANT_HEADERS = [
    "product_id",
    "legacy_id",
    "name",
    "slug",
    "active",
    "status",
    "price",
    "regular_price",
    "short_description",
    "description",
    "image",
    "gallery",
    "product_type",
    "collection",
    "tags",
    "featured",
    "sort_order",
    "whatsapp_enabled",
    "seo_title",
    "seo_description",
]

POT_HEADERS = [
    "product_id",
    "model_name",
    "product_slug",
    "collection",
    "material",
    "size",
    "colour",
    "price",
    "regular_price",
    "active",
    "status",
    "description",
    "short_description",
    "image",
    "gallery",
    "featured",
    "sort_order",
    "whatsapp_enabled",
    "seo_title",
    "seo_description",
    "legacy_slug",
]

NON_PLANT = {"Soil & Fertilizer", "Stone's"}
POT_ONLY = {"Pots"}


def infer_tags(categories):
    tags = []
    cats = set(categories)
    mapping = {
        "Indoor Plants": "Indoor",
        "Outdoor Plants": "Outdoor",
        "Succulents": "Succulents",
        "Gifting Plants": "Gifting",
        "Semi Indoor Plants": "Indoor",
        "Seasonal Plants": "Flowering",
        "Landscaping Plants": "Outdoor",
        "Creepers/Hanging": "Outdoor",
    }
    for c, t in mapping.items():
        if c in cats and t not in tags:
            tags.append(t)
    if "Indoor Plants" in cats or "Succulents" in cats:
        if "Low Maintenance" not in tags:
            tags.append("Low Maintenance")
    if any(x in cats for x in ("Indoor Plants", "Semi Indoor Plants")):
        if "Air Purifying" not in tags:
            tags.append("Air Purifying")
    return ", ".join(tags)


def infer_pot_collection(name, slug):
    n = (name or "").lower()
    s = slug or ""
    if "eco" in n or "eco" in s:
        return "ECO SERIES"
    if "illumin" in n or "led" in n:
        return "Illuminated & Decorative"
    if "hang" in n or "hook" in n:
        return "Hanging Planters"
    if any(k in n for k in ("milano", "tokyo", "paris", "venice", "rome", "valencia")):
        return "Statement Planters"
    return "Plastic Pots"


def infer_pot_model(slug, name):
    s = slug or ""
    if s.startswith("verona-"):
        return "VERONA ECO", "verona-eco"
    m = re.match(r"^(bello-square-20)", s)
    if m:
        return "BELLO SQUARE 20", m.group(1)
    m = re.match(r"^(valencia-\d+)", s)
    if m:
        return f"VALENCIA {m.group(1).split('-')[-1]}", "valencia-planters"
    if s.startswith("valencia"):
        return "VALENCIA PLANTERS", "valencia-planters"
    base = re.sub(r"-set-of-\d+.*$", "", s)
    base = re.sub(r"-\d+$", "", base)
    model_name = (name or slug).upper()
    return model_name[:80], slugify(base or s)


def slugify(s):
    s = (s or "").lower().replace("&", "and").replace("'", "")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")


def infer_size(slug, name):
    m = re.search(r"verona-(\d+(?:\.\d+)?)", slug or "")
    if m:
        return m.group(1)
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:inch|cm|\"|$)", (name or "").lower())
    if m:
        return m.group(1)
    m = re.search(r"-(\d+(?:\.\d+)?)(?:-|$)", slug or "")
    if m:
        return m.group(1)
    return "Standard"


def is_plant_row(p):
    cats = set(p.get("categories") or [])
    if cats & NON_PLANT:
        return False
    if cats <= POT_ONLY:
        return False
    if "Pots" in cats and len(cats) == 1:
        return False
    return True


def is_pot_row(p):
    cats = set(p.get("categories") or [])
    return bool(cats & POT_ONLY) or (
        "Pots" in cats and not (cats - POT_ONLY - {"Plants with Pots"})
    )


def main():
    catalog = json.load(open(CATALOG_PATH, encoding="utf-8"))
    os.makedirs(os.path.dirname(os.path.join(ROOT, EXCEL_PATH)), exist_ok=True)
    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "Plants"
    ws_p.append(PLANT_HEADERS)
    for c in PLANT_HEADERS:
        ws_p.cell(row=1, column=PLANT_HEADERS.index(c) + 1).font = Font(bold=True)

    sort = 0
    for p in catalog:
        if not is_plant_row(p):
            continue
        sort += 1
        slug = p["slug"]
        active = slug in DEFAULT_ACTIVE_PLANT_SLUGS
        ws_p.append(
            [
                f"P-{p.get('id', slug)}",
                p.get("id", ""),
                p.get("name", ""),
                slug,
                "YES" if active else "NO",
                "active" if active else "archived",
                p.get("price", ""),
                p.get("regular_price", ""),
                (p.get("short_description") or "")[:500],
                (p.get("description") or "")[:2000],
                p.get("image", ""),
                "|".join(p.get("gallery") or []),
                "plant",
                "",
                infer_tags(p.get("categories") or []),
                "YES" if active else "NO",
                sort,
                "YES",
                f"{p.get('name', '')} | Indore Nursery",
                (p.get("short_description") or p.get("name", ""))[:155],
            ]
        )

    ws_o = wb.create_sheet("Pots")
    ws_o.append(POT_HEADERS)
    for c in POT_HEADERS:
        ws_o.cell(row=1, column=POT_HEADERS.index(c) + 1).font = Font(bold=True)

    sort = 0
    seen_pot_slugs = set()
    for p in catalog:
        cats = set(p.get("categories") or [])
        if "Pots" not in cats:
            continue
        if cats <= NON_PLANT:
            continue
        slug = p["slug"]
        if "set-of" in slug and not slug.startswith("verona-"):
            continue
        model_name, product_slug = infer_pot_model(slug, p.get("name"))
        sort += 1
        pid = f"PT-{p.get('id', slug)}"
        ws_o.append(
            [
                pid,
                model_name,
                product_slug,
                infer_pot_collection(p.get("name", ""), slug),
                "Plastic",
                infer_size(slug, p.get("name")),
                "Standard",
                p.get("price", ""),
                p.get("regular_price", ""),
                "YES",
                "active",
                (p.get("description") or "")[:2000],
                (p.get("short_description") or "")[:500],
                p.get("image", ""),
                "|".join(p.get("gallery") or []),
                "NO",
                sort,
                "YES",
                f"{model_name} Planters | Indore Nursery",
                (p.get("short_description") or model_name)[:155],
                slug,
            ]
        )
        seen_pot_slugs.add(slug)

    out = os.path.join(ROOT, EXCEL_PATH)
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Plants rows: {ws_p.max_row - 1}, Pot variant rows: {ws_o.max_row - 1}")


if __name__ == "__main__":
    main()
