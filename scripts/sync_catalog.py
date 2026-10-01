#!/usr/bin/env python3
"""Sync product catalog from Google Sheets or local Excel → data/generated/*.json."""

from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)


def _load_dotenv() -> None:
    path = os.path.join(ROOT, ".env")
    if not os.path.isfile(path):
        return
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, val = line.split("=", 1)
            key = key.strip()
            val = val.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = val


from lib.catalog_sync import SyncError, rows_from_table, sync_from_row_dicts
from lib.config import EXCEL_PATH
from lib.google_sheets_catalog import load_catalog_rows_from_sheets, sheets_configured


def load_from_excel() -> tuple[list, list, str]:
    from openpyxl import load_workbook

    path = os.path.join(ROOT, EXCEL_PATH)
    if not os.path.isfile(path):
        raise SyncError(
            f"Excel not found at {path}. Run: python scripts/migrate_catalog_to_excel.py"
        )

    def load_sheet(wb, name):
        if name not in wb.sheetnames:
            raise SyncError(f"Missing sheet: {name}")
        ws = wb[name]
        table = list(ws.iter_rows(values_only=True))
        return rows_from_table([list(r) for r in table])

    wb = load_workbook(path, read_only=True, data_only=True)
    plants = load_sheet(wb, "Plants")
    pots = load_sheet(wb, "Pots")
    return plants, pots, "excel"


def main() -> int:
    _load_dotenv()
    if sheets_configured():
        plant_rows, pot_rows, source = load_catalog_rows_from_sheets()
        label = "Google Sheets"
    else:
        plant_rows, pot_rows, source = load_from_excel()
        label = "Excel"

    n_plants, n_pots = sync_from_row_dicts(plant_rows, pot_rows, source=source)
    print(f"Synced {n_plants} plants, {n_pots} pot models from {label} → data/generated/")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SyncError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        raise SystemExit(1)
