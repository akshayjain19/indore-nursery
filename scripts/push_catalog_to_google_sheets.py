#!/usr/bin/env python3
"""One-time: upload local Excel Plants/Pots tabs to a Google Spreadsheet."""

from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import load_workbook

from lib.catalog_sync import SyncError
from lib.config import EXCEL_PATH
from lib.google_sheets_catalog import authorize_client, service_account_email


def main() -> int:
    spreadsheet_id = os.environ.get("GOOGLE_SHEETS_SPREADSHEET_ID", "").strip()
    if not spreadsheet_id:
        print("Set GOOGLE_SHEETS_SPREADSHEET_ID to the target spreadsheet.", file=sys.stderr)
        return 1

    path = os.path.join(ROOT, EXCEL_PATH)
    if not os.path.isfile(path):
        print(f"Missing {path}. Run migrate_catalog_to_excel.py first.", file=sys.stderr)
        return 1

    wb = load_workbook(path, read_only=True, data_only=True)
    gc = authorize_client(write=True)
    sh = gc.open_by_key(spreadsheet_id)

    for name in ("Plants", "Pots"):
        if name not in wb.sheetnames:
            raise SyncError(f"Workbook missing sheet {name}")
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        try:
            ws = sh.worksheet(name)
        except Exception:
            ws = sh.add_worksheet(title=name, rows=len(rows) + 1, cols=max(len(r) for r in rows))
        ws.clear()
        ws.update(rows, value_input_option="USER_ENTERED")
        print(f"Uploaded {len(rows) - 1} data rows → tab {name!r}")

    email = service_account_email()
    print(f"\nShare the spreadsheet with this service account (Editor): {email}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SyncError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        raise SystemExit(1)
